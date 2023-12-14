# -*- coding: utf-8 -*-
"""
Created on Mon Nov 28 09:07:30 2022

@author: aborst
"""

import docx
import numpy as np
import matplotlib.pyplot as plt
import sort_library_NEW as sl

direct = 'SIPaper_Data/OpticLobeDM/'

fname1='central_column_connectivity.csv'
fname2='offset_column_connectivity.csv'
    
nofcells=65

save_switch=0

def setmyaxes(myxpos,myypos,myxsize,myysize):
    
    ax=plt.axes([myxpos,myypos,myxsize,myysize])
    ax.xaxis.set_ticks_position('bottom')
    ax.yaxis.set_ticks_position('left')

def read_single_connM(fname):
    
    print('reading data from ' + fname + ' file')
    
    data=np.genfromtxt(direct+fname, delimiter=',', dtype=None)
    
    connM=np.zeros((nofcells,nofcells))
    ctype = ["" for x in range(nofcells)]
    
    for i in range(nofcells):
        ctype[i]=data[i+1,0].decode()
    
    for i in range(1,nofcells+1,1):
        for j in range(1,nofcells+1,1):
            if data[i,j] == b'':
                connM[i-1,j-1] = 0
            else:
                connM[i-1,j-1]=float(data[i,j])
            
    connM=np.transpose(connM)
    
    return connM, ctype

rawM, ctype = read_single_connM(fname1)
ctype = np.array(ctype)

mydim=65
counter=0
thr=4

z=np.arange(mydim)
z_bounds=[(0,mydim)]
for i in range(mydim-1): z_bounds.append((0,mydim))

Mzi=np.zeros((mydim,mydim))
Mzj=np.zeros((mydim,mydim))

fontsize_title=10

def calcAdjM(M,thr=0,mode='both'):
    
    mydim=M.shape[0]

    AdjM=np.zeros((mydim,mydim))
    
    if mode=='both':
    
        AdjM=1*(abs(M)>thr)
                    
    if mode=='exc':
    
        AdjM=1*(M>thr)
        
    if mode=='inh':
    
        AdjM=1*(M<(-thr))
                
    return AdjM 

def plot_nofconnections():
    
    nofconnections=np.zeros((3,100))
    
    mylw=2
    
    for i in range(100):
        
        nofconnections[0,i]=np.sum(calcAdjM(rawM,thr=i,mode='both'))
        nofconnections[1,i]=np.sum(calcAdjM(rawM,thr=i,mode='exc'))
        nofconnections[2,i]=np.sum(calcAdjM(rawM,thr=i,mode='inh'))
        
    plt.plot(nofconnections[0],linewidth=mylw,color='black',label='all')
    plt.plot(nofconnections[1],linewidth=mylw,color='blue',label='exc')
    plt.plot(nofconnections[2],linewidth=mylw,color='red',label='inh')
    plt.yscale('log')
    plt.xscale('log')
    plt.ylim(1,1000)
    plt.xlim(1,100)
    plt.plot([thr,thr],[1,1000],linestyle='dashed',color='black')
    plt.xlabel('threshold')
    plt.ylabel('total number of connections')
    plt.legend(loc=1,frameon=False)
    plt.title('connection number = f(thrld)',fontsize=fontsize_title)
        
M=calcAdjM(rawM, thr=thr)

print('total number connections > '+ str(thr) + ' synapses = '+ str(int(np.sum(M))))

def plotM(M,title):
    
    plt.imshow(M)
    plt.axis('off')
    plt.title(title,fontsize=fontsize_title)   
    
def plot_all(M1,M2):
    
    plt.figure(figsize=(7,3))
    plt.subplot(1,2,1)
    plotM(M1,'original')
    plt.subplot(1,2,2)
    plotM(M2,'rc minimized')
    
def plot_full_M(M,mylabels,mytitle,maxval=1):
  
    plt.imshow(M,vmin=-maxval,vmax=maxval,cmap='coolwarm',interpolation='None') 
    plt.xticks(np.arange(mydim), mylabels, rotation=90, fontsize=4)
    plt.yticks(np.arange(mydim), mylabels, rotation=00, fontsize=4)
    plt.title(mytitle, fontsize=fontsize_title)

def calc_rc_number(M):
    
    rec_number=int(np.sum(np.triu(M)))
    
    return rec_number

def go(plotit=1):
    
    rc_number=calc_rc_number(M)
    print('original M, number of recurrent synapses:  ', rc_number)
    
    srtM = sl.SI_sort(M)
            
    if plotit==1: plot_all(M,srtM)
    
    rc_number=calc_rc_number(srtM)
    print('reordered M, number of recurrent synapses:  ', rc_number)
    
    return srtM
    
def identify_rec_synapses(n, trld = 0.2):
    
    rec_synM  = np.zeros((mydim,mydim))
    
    least_nofrecsyn = np.sum(np.triu(M,1))
    
    for i in range(n):
        
        srtM,argsort_z=sl.SI_sort(M,return_object='both')
        recM=np.triu(srtM,1)
        
        print('Run ', i, ': nof rec synpases = ', np.sum(np.triu(srtM,1)))
        
        rev_z=sl.reverse_permutation(argsort_z)
        rec_synM += sl.renumberM(recM,rev_z)
        
        if np.sum(np.triu(srtM,1)) < least_nofrecsyn:
            
            least_nofrecsyn = np.sum(np.triu(srtM,1))
            best_argsort_z  = argsort_z
              
    rec_synM = rec_synM/(1.0*n)
    
    plt.imshow(rec_synM * (rec_synM > trld),cmap='Reds')
    plt.xticks(np.arange(mydim), ctype, rotation=90, fontsize=6)
    plt.yticks(np.arange(mydim), ctype, rotation=00, fontsize=6)
    plt.title('recurrent synapses', fontsize=fontsize_title)
    
    cbar = plt.colorbar()
    cbar.set_label('probability', rotation=90, fontsize=10)

    np.save('IntraCol_rec_synM.npy',rec_synM)
    np.save('IntraCol_best_argsort_z.npy',best_argsort_z)
    
    return rec_synM

def plot_figure():
    
    global thr
    
    thr=4
    
    argsort_z=np.load(direct+'IntraCol_best_argsort_z.npy')
    
    newM_adj=sl.renumberM(M,argsort_z)
    newM_raw=sl.renumberM(rawM,argsort_z)
    newlabels=ctype[argsort_z.astype(int)]
    
    plt.figure(figsize=(7.5,10))
    
    setmyaxes(0.06,0.60,0.4,0.3)
    plot_full_M(rawM,ctype,'original M', maxval=30)
    
    setmyaxes(0.63,0.70,0.32,0.2)
    plot_nofconnections()
    
    setmyaxes(0.70,0.38,0.25,0.25)
    plotM(M,'original adjacency M')
    
    setmyaxes(0.70,0.15,0.25,0.25)
    plotM(newM_adj,'reordered adjacency M')
    
    setmyaxes(0.065,0.20,0.5,0.3)
    plot_full_M(newM_raw,newlabels,'reordered M', maxval=30)
    cbar = plt.colorbar()
    cbar.set_ticks([-30,-20,-10,0,10,20,30])
    cbar.set_ticklabels(['<30','20','10','0','10','20','>30'])
    cbar.set_label('inhibitory         # of synapses       excitatory', rotation=90, fontsize=10)
    
    nof_rcc_M=np.sum(np.triu(M))
    nof_rcc_newM=np.sum(np.triu(newM_adj))
    
    print('Number of recurrent connections, original M = ' + str(nof_rcc_M))
    print('Number of recurrent connections, reordered M = ' + str(nof_rcc_newM))
    
    if save_switch == 1:
        
        plt.savefig('Figure_6.tiff',dpi=300)
    

def plot_table(trld=0.5):

    rec_synM = np.load(direct+'IntraCol_rec_synM.npy')
    rec_synM = rec_synM * (rec_synM > trld)
    
    yval, xval = np.where(rec_synM != 0)
    
    nofsyns = yval.size
    
    # sort along p_value
    
    myarray = np.zeros(nofsyns)
    
    for i in range(nofsyns):
        
        myarray[i] = rec_synM[yval[i],xval[i]]
        
    newargs = np.argsort(myarray)
    newargs = newargs[::-1]
    
    mytable=[['a','b','c','d']]
    for i in range(nofsyns-1):
        mytable.append(['a', 'b', 'c', 'd'])
    
    print('p_recu   syn_str   connection')
    
    # now use newargs as index for table
    
    for k in range(nofsyns):
        
        i = newargs[k]
        
        mytable[k][0] = str(format(rec_synM[yval[i],xval[i]],'.3f'))
        mytable[k][1] = str(int(rawM[yval[i],xval[i]]))
        mytable[k][2] = ctype[xval[i]]
        mytable[k][3] = ctype[yval[i]]
    
    fig,ax1 = plt.subplots(figsize=(7.5,10))
    
    myrows = np.arange(nofsyns)+1
    mycols = ['p_recu','syn_str','pre','post']
    
    plt.axis('off')
    
    table = ax1.table(cellText   = mytable, 
                      rowLabels  = myrows, 
                      colLabels  = mycols,
                      rowColours = ['palegreen'] * nofsyns,
                      colColours = ['palegreen'] * nofsyns,
                      cellLoc    = 'center', 
                      loc        = 'center')
    
    table.scale(0.6,1)
    table.set_fontsize(8)
    fig.tight_layout()
    
    if save_switch == 1:
        
        plt.savefig('Table_1.tiff',dpi=300)
        
def export_table_to_word(trld=0.5):

    rec_synM = np.load(direct+'IntraCol_rec_synM.npy')
    rec_synM = rec_synM * (rec_synM > trld)
    
    yval, xval = np.where(rec_synM != 0)
    
    nofsyns = yval.size
    
    # sort along p_value
    
    myarray = np.zeros(nofsyns)
    
    for i in range(nofsyns):
        
        myarray[i] = rec_synM[yval[i],xval[i]]
        
    newargs = np.argsort(myarray)
    newargs = newargs[::-1]
    
    mytable=[['a','b','c','d']]
    for i in range(nofsyns-1):
        mytable.append(['a', 'b', 'c', 'd'])
    
    print('p_recu   syn_str   connection')
    
    # now use newargs as index for table
    
    for k in range(nofsyns):
        
        i = newargs[k]
        
        mytable[k][0] = str(format(rec_synM[yval[i],xval[i]],'.3f'))
        mytable[k][1] = str(int(rawM[yval[i],xval[i]]))
        mytable[k][2] = ctype[xval[i]]
        mytable[k][3] = ctype[yval[i]]
        
    # Initialise the Word document
    doc = docx.Document()
    
    # Initialise the table
    t = doc.add_table(rows=nofsyns, cols=4)
    
    for i in range(nofsyns):
        for j in range(4):
            cell = mytable[i][j]
            t.cell(i, j).text = str(cell)
    
    myrows = np.arange(nofsyns)+1
    mycols = ['p_recu','syn_str','pre','post']
        
    doc.save('table 1.docx')
    


        
    
                
    
    
            

    

            

            
            
        